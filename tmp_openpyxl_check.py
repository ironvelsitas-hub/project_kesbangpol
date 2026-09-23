from openpyxl import load_workbook

wb = load_workbook(r'C:\Users\LENOVO\Downloads\Projek web kesbangpol\Data Lokasi Tawuran 2026 JAKARTA BARAT.xlsx', data_only=True)
print(wb.sheetnames)
ws = wb[wb.sheetnames[0]]
for row in ws.iter_rows(min_row=1, max_row=12, values_only=True):
    print(row[:9])
